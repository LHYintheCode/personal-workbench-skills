#!/usr/bin/env python3
"""Inventory Douyin Creator Center XLSX exports without changing the originals."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Create a workbook manifest, field catalog, previews, and CSV mirrors."
    )
    parser.add_argument("--input", required=True, type=Path, help="Directory containing XLSX exports")
    parser.add_argument("--output", required=True, type=Path, help="Directory for derived inventory files")
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def serializable(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def text_value(value: Any) -> str:
    if value is None:
        return ""
    value = serializable(value)
    return str(value)


def safe_name(value: str) -> str:
    value = value.replace("/", "__").replace("\\", "__")
    value = re.sub(r"[\x00-\x1f:*?\"<>|]", "_", value)
    return re.sub(r"\s+", "_", value).strip("._") or "sheet"


def normalize_headers(values: Iterable[Any], width: int) -> list[str]:
    headers: list[str] = []
    seen: Counter[str] = Counter()
    raw = list(values)
    for index in range(width):
        base = text_value(raw[index] if index < len(raw) else "").strip() or f"未命名字段_{index + 1}"
        seen[base] += 1
        headers.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return headers


def classify(value: Any) -> str:
    if value is None or value == "":
        return "empty"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "numeric"
    if isinstance(value, (datetime, date, time)):
        return "datetime"
    return "text"


def write_csv(path: Path, headers: list[str], rows: list[list[Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        for row in rows:
            writer.writerow([text_value(value) for value in row])


def main() -> int:
    args = parse_args()
    input_dir = args.input.expanduser().resolve()
    output_dir = args.output.expanduser().resolve()
    if not input_dir.is_dir():
        raise SystemExit(f"Input directory does not exist: {input_dir}")

    workbooks = sorted(input_dir.rglob("*.xlsx"))
    if not workbooks:
        raise SystemExit(f"No XLSX files found under: {input_dir}")

    output_dir.mkdir(parents=True, exist_ok=True)
    sheets_dir = output_dir / "sheets"
    manifest_rows: list[dict[str, Any]] = []
    workbook_inventory: list[dict[str, Any]] = []
    field_rows: list[dict[str, Any]] = []
    total_sheets = 0
    total_data_rows = 0

    for workbook_path in workbooks:
        relative = workbook_path.relative_to(input_dir)
        stat = workbook_path.stat()
        workbook_hash = sha256(workbook_path)
        workbook_entry: dict[str, Any] = {
            "relative_path": relative.as_posix(),
            "sha256": workbook_hash,
            "size_bytes": stat.st_size,
            "modified_at": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
            "sheets": [],
        }

        workbook = load_workbook(
            workbook_path,
            data_only=False,
            read_only=True,
            keep_links=False,
        )
        for sheet in workbook.worksheets:
            total_sheets += 1
            all_rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            width = max((len(row) for row in all_rows), default=0)
            header_source = all_rows[0] if all_rows else []
            headers = normalize_headers(header_source, width)
            data_rows = [
                (row + [None] * (width - len(row)))[:width]
                for row in all_rows[1:]
            ]
            total_data_rows += len(data_rows)

            csv_name = safe_name(f"{relative.with_suffix('').as_posix()}__{sheet.title}") + ".csv"
            csv_relative = Path("sheets") / csv_name
            write_csv(output_dir / csv_relative, headers, data_rows)

            field_stats: list[dict[str, Any]] = []
            for column_index, header in enumerate(headers):
                values = [row[column_index] for row in data_rows]
                type_counts = Counter(classify(value) for value in values)
                distinct = []
                seen_values: set[str] = set()
                for value in values:
                    rendered = text_value(value)
                    if rendered and rendered not in seen_values:
                        seen_values.add(rendered)
                        distinct.append(rendered)
                    if len(distinct) >= 5:
                        break
                field_entry = {
                    "workbook": relative.as_posix(),
                    "sheet": sheet.title,
                    "field": header,
                    "data_rows": len(data_rows),
                    "non_empty_count": len(values) - type_counts["empty"],
                    "missing_count": type_counts["empty"],
                    "distinct_count": len({text_value(value) for value in values if text_value(value)}),
                    "numeric_count": type_counts["numeric"],
                    "datetime_count": type_counts["datetime"],
                    "text_count": type_counts["text"],
                    "boolean_count": type_counts["boolean"],
                    "sample_values": " | ".join(distinct),
                }
                field_rows.append(field_entry)
                field_stats.append(field_entry)

            preview = [
                {headers[index]: serializable(row[index]) for index in range(width)}
                for row in data_rows[:3]
            ]
            workbook_entry["sheets"].append(
                {
                    "title": sheet.title,
                    "row_count_including_header": len(all_rows),
                    "data_row_count": len(data_rows),
                    "column_count": width,
                    "headers": headers,
                    "csv_mirror": csv_relative.as_posix(),
                    "preview": preview,
                    "fields": field_stats,
                }
            )

        workbook.close()
        workbook_inventory.append(workbook_entry)
        manifest_rows.append(
            {
                "relative_path": relative.as_posix(),
                "sha256": workbook_hash,
                "size_bytes": stat.st_size,
                "modified_at": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
                "sheet_count": len(workbook_entry["sheets"]),
                "data_row_count": sum(sheet["data_row_count"] for sheet in workbook_entry["sheets"]),
            }
        )

    inventory_payload = {
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "input_dir": str(input_dir),
        "workbook_count": len(workbook_inventory),
        "sheet_count": total_sheets,
        "data_row_count": total_data_rows,
        "workbooks": workbook_inventory,
    }
    (output_dir / "workbook_inventory.json").write_text(
        json.dumps(inventory_payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    manifest_headers = [
        "relative_path",
        "sha256",
        "size_bytes",
        "modified_at",
        "sheet_count",
        "data_row_count",
    ]
    with (output_dir / "source_files.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=manifest_headers)
        writer.writeheader()
        writer.writerows(manifest_rows)

    field_headers = [
        "workbook",
        "sheet",
        "field",
        "data_rows",
        "non_empty_count",
        "missing_count",
        "distinct_count",
        "numeric_count",
        "datetime_count",
        "text_count",
        "boolean_count",
        "sample_values",
    ]
    with (output_dir / "field_catalog.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=field_headers)
        writer.writeheader()
        writer.writerows(field_rows)

    summary = [
        "# 抖音创作者中心导出文件字段清单",
        "",
        f"- 生成时间：{inventory_payload['generated_at']}",
        f"- 原始目录：`{input_dir}`",
        f"- Excel 文件：{len(workbook_inventory)} 个",
        f"- 工作表：{total_sheets} 个",
        f"- 数据行：{total_data_rows} 行（不含表头）",
        f"- 去重后的字段名：{len({row['field'] for row in field_rows})} 个",
        "",
        "## 输出",
        "",
        "- `source_files.csv`：原始文件哈希、大小、工作表数和数据行数。",
        "- `field_catalog.csv`：每个工作表的字段、缺失数、类型和样例值。",
        "- `workbook_inventory.json`：完整机器可读清单与每表前三行预览。",
        "- `sheets/`：每个工作表的 UTF-8 CSV 镜像；不修改 Excel 原件。",
        "",
        "## 读取原则",
        "",
        "- 空值保持为空，不把缺失值改写为 0。",
        "- 百分比和带单位值保留官方导出的原始表示。",
        "- 同一指标在不同导出时刻可能不同，以文件时间和 SHA-256 区分。",
    ]
    (output_dir / "README.md").write_text("\n".join(summary) + "\n", encoding="utf-8")

    print(
        json.dumps(
            {
                "workbook_count": len(workbook_inventory),
                "sheet_count": total_sheets,
                "data_row_count": total_data_rows,
                "field_occurrence_count": len(field_rows),
                "distinct_field_count": len({row["field"] for row in field_rows}),
                "output_dir": str(output_dir),
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
